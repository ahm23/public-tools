"""
Convert a numbered-list DOCX to XLSX with columns: Title, Type, Text.

Usage:
    python req-docx_to_xlsx.py thing.docx output.xlsx
"""

import sys
from docx import Document
from docx.oxml.ns import qn
from openpyxl import Workbook


def get_num_info(paragraph):
    """Return (numId, ilvl) for a paragraph, or (None, None) if not numbered."""
    pPr = paragraph._element.find(qn('w:pPr'))
    if pPr is None:
        return None, None
    numPr = pPr.find(qn('w:numPr'))
    if numPr is None:
        return None, None
    ilvl_el = numPr.find(qn('w:ilvl'))
    numId_el = numPr.find(qn('w:numId'))
    ilvl = int(ilvl_el.get(qn('w:val'))) if ilvl_el is not None else None
    numId = int(numId_el.get(qn('w:val'))) if numId_el is not None else None
    return numId, ilvl


def has_drawing(paragraph):
    """Check if a paragraph contains an image/drawing."""
    ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    return len(paragraph._element.findall(f'.//{ns}drawing')) > 0


def get_bold_text(paragraph):
    """Return the bold portion of a paragraph's text.
    
    Checks both explicit run-level bold and paragraph-style bold."""
    # Check paragraph style first
    style_bold = False
    if paragraph.style and paragraph.style.font:
        style_bold = paragraph.style.font.bold is True
    
    # Collect runs with explicit bold, or all runs if style is bold
    parts = []
    for r in paragraph.runs:
        is_bold = (r.bold is True) or (r.font.bold is True) or style_bold
        if is_bold:
            parts.append(r.text)
    
    return ''.join(parts).strip()


def is_figure_or_table_caption(text, bold_text):
    """Return True if this is a Figure/Table caption to skip."""
    if not bold_text:
        return False
    t = text.strip()
    return t.startswith('Figure') or t.startswith('Table')


def get_clean_text(paragraph):
    """Get paragraph text, replacing embedded images with '[IMAGE]'."""
    text = ''.join(r.text for r in paragraph.runs).strip()
    if has_drawing(paragraph):
        if text:
            text += " [IMAGE]"
        else:
            text = "[IMAGE]"
    return text


def convert_docx_to_xlsx(docx_path, xlsx_path):
    """Main conversion function."""
    doc = Document(docx_path)
    
    rows = []
    current_title = ""
    current_body_parts = []
    
    def flush_current():
        nonlocal current_title, current_body_parts
        if current_title:
            body = "\n".join(current_body_parts).strip()
            rows.append((current_title, body))
        current_title = ""
        current_body_parts = []
    
    for para in doc.paragraphs:
        numId, ilvl = get_num_info(para)
        bold_text = get_bold_text(para)
        clean_text = get_clean_text(para)
        
        # Skip empty paragraphs
        if not clean_text and not has_drawing(para) and numId is None:
            continue
        
        # Skip figure/table captions entirely
        if is_figure_or_table_caption(clean_text, bold_text):
            continue
        
        # Classify this paragraph
        is_header = numId is not None and bool(bold_text)
        
        if is_header:
            flush_current()
            current_title = clean_text
        
        elif numId is not None:
            # Numbered but not bold → bullet (any numId)
            current_body_parts.append(f"- {clean_text}")
        
        elif numId is None and clean_text:
            # Plain body text
            current_body_parts.append(clean_text)
        
        elif has_drawing(para):
            # Image paragraph
            current_body_parts.append(clean_text)
        
        elif clean_text:
            # Catch-all
            current_body_parts.append(clean_text)
    
    flush_current()
    
    # Write XLSX
    wb = Workbook()
    ws = wb.active
    ws.title = "Specification"
    ws.append(["Title", "Type", "Text"])
    
    for title, text in rows:
        row_type = "Heading" if not text.strip() else "Requirement"
        ws.append([title, row_type, text])
    
    ws.column_dimensions['A'].width = 40
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 80
    wb.save(xlsx_path)
    
    print(f"Written {len(rows)} rows to {xlsx_path}")
    return rows


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    
    docx_path = sys.argv[1]
    xlsx_path = sys.argv[2] if len(sys.argv) > 2 else docx_path.replace('.docx', '.xlsx')
    
    print(f"Reading {docx_path}...")
    rows = convert_docx_to_xlsx(docx_path, xlsx_path)
    
    if not rows:
        # Diagnostic info: show what types of paragraphs were found
        print("\nNo rows produced. Diagnostic info:")
        doc = Document(docx_path)
        for i, p in enumerate(doc.paragraphs[:30]):
            text = ''.join(r.text for r in p.runs)[:80]
            if not text:
                text = "(empty)"
            numId, ilvl = get_num_info(p)
            bt = get_bold_text(p)
            print(f"  [{i:2d}] numId={str(numId):4s} ilvl={str(ilvl):4s} bold={repr(bt[:30]):35s} text={repr(text[:60])}")
    
    print("\nPreview:")
    for i, (title, body) in enumerate(rows):
        row_type = "Heading" if not body.strip() else "Requirement"
        bp = body[:80].replace('\n', '\\n')
        print(f"  {i+1:2d}. [{row_type:11s}] {title:40s} | {bp}")
